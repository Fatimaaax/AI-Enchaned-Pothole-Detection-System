import os
import cv2
import tkinter as tk
from tkinter import filedialog
from PIL import Image, ImageTk
from ultralytics import YOLO
import serial
import serial.tools.list_ports
from datetime import datetime
import pandas as pd
from openpyxl import load_workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.utils import get_column_letter

last_processed_frame = None

try:
    model = YOLO("/home/fatima/Desktop/Road/run_1/best.pt")
except Exception as e:
    print(f"Failed to load model: {e}")
    exit(1)

detection_threshold = 0.5
cap = None
running = False
ser = None
gps_data = "N/A"
log_dir = "temp_logs"
os.makedirs(log_dir, exist_ok=True)
log_file = os.path.join(log_dir, "detection_log.xlsx")

if not os.path.isfile(log_file):
    df = pd.DataFrame(columns=["Timestamp", "Defect Type", "GPS Data", "Image", "Image Path"])
    df.to_excel(log_file, index=False)

root = tk.Tk()
root.title("Road Damage Detection - YOLOv8")
root.geometry("1100x650")

control_frame = tk.Frame(root)
control_frame.pack(side=tk.LEFT, fill=tk.Y, padx=10, pady=10)

panel = tk.Label(root)
panel.pack(side=tk.LEFT, padx=10, pady=10)

textbox = tk.Text(root, height=30, width=30)
textbox.pack(side=tk.RIGHT, padx=10, pady=10)
textbox.insert(tk.END, "Detected Defects:\n")

def list_serial_ports():
    ports = serial.tools.list_ports.comports()
    return [port.device for port in ports]

def connect_serial():
    global ser
    try:
        port = com_port_var.get()
        baud = int(baud_rate_var.get())
        ser = serial.Serial(port, baud, timeout=1)
        print(f"Connected to {port} at {baud} baud.")
        connect_btn.config(bg="#39FF14")
    except Exception as e:
        print(f"Failed to connect: {e}")
        ser = None
        connect_btn.config(bg="SystemButtonFace")

def disconnect_serial():
    global ser
    if ser and ser.is_open:
        ser.close()
        print("Serial port disconnected.")
    ser = None
    connect_btn.config(bg="SystemButtonFace")

def read_gps_data():
    global gps_data
    if ser and ser.is_open:
        try:
            line = ser.readline().decode(errors='ignore').strip()
            if line.startswith("<") and line.endswith(">"):
                gps_data = line
        except Exception:
            gps_data = "N/A"

def update_gps_label():
    read_gps_data()
    gps_label.config(text=f"GPS: {gps_data}")
    root.after(1000, update_gps_label)

def update_threshold(value):
    global detection_threshold
    detection_threshold = float(value)

def show_image(frame):
    rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
    img = Image.fromarray(rgb)
    imgtk = ImageTk.PhotoImage(image=img)
    panel.imgtk = imgtk
    panel.config(image=imgtk)

def log_detection(defect_type, image_with_boxes):
    timestamp = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    read_gps_data()
    filename = f"{timestamp.replace(':', '-')}.jpg"
    filepath = os.path.join(log_dir, filename)

    resized = cv2.resize(image_with_boxes, (100, 100))
    cv2.imwrite(filepath, resized)

    df = pd.read_excel(log_file)
    new_row = {
        "Timestamp": timestamp,
        "Defect Type": defect_type,
        "GPS Data": gps_data,
        "Image": "",
        "Image Path": os.path.relpath(filepath)
    }
    df = pd.concat([df, pd.DataFrame([new_row])], ignore_index=True)
    df.to_excel(log_file, index=False)

def save_images_to_excel():
    wb = load_workbook(log_file)
    ws = wb.active
    df = pd.read_excel(log_file)

    for col, width in {"A": 20, "B": 20, "C": 30, "D": 18, "E": 40}.items():
        ws.column_dimensions[col].width = width

    for idx, row in df.iterrows():
        img_path = row["Image Path"]
        row_num = idx + 2
        if os.path.exists(img_path):
            try:
                img = ExcelImage(img_path)
                img.width, img.height = 100, 100
                ws.row_dimensions[row_num].height = 80
                ws.add_image(img, f"D{row_num}")
            except Exception as e:
                print(f"Failed to add image {img_path}: {e}")

    wb.save(log_file)
    print("Images embedded and sheet formatted.")

def display_defects(results, img=None):
    textbox.delete('1.0', tk.END)
    textbox.insert(tk.END, "Detected Defects:\n")
    logged_classes = set()

    for r in results:
        if r.boxes is not None:
            for box in r.boxes:
                cls_id = int(box.cls[0])
                conf = float(box.conf[0])
                class_name = model.names[cls_id]
                textbox.insert(tk.END, f"{class_name} ({conf:.2f})\n")
                if img is not None and class_name not in logged_classes:
                    log_detection(class_name, img)
                    logged_classes.add(class_name)

def open_image():
    file_path = filedialog.askopenfilename(filetypes=[("Image files", "*.jpg *.jpeg *.png")])
    if file_path:
        image = cv2.imread(file_path)
        results = model(image, conf=detection_threshold)
        frame_copy = image.copy()

        for r in results:
            if r.boxes is not None:
                for box in r.boxes:
                    x1, y1, x2, y2 = map(int, box.xyxy[0])
                    cls_id = int(box.cls[0])
                    conf = float(box.conf[0])
                    label = f"{model.names[cls_id]} ({conf:.2f})"
                    cv2.rectangle(frame_copy, (x1, y1), (x2, y2), (57, 255, 20), 2)
                    cv2.putText(frame_copy, label, (x1, y1 - 10), cv2.FONT_HERSHEY_SIMPLEX, 0.6, (0, 255, 255), 2)

        show_image(frame_copy)
        display_defects(results, img=frame_copy)

def open_video():
    file_path = filedialog.askopenfilename(filetypes=[("Video files", "*.mp4 *.avi *.mov")])
    if file_path:
        start_video(file_path)

def start_camera():
    start_video(0)
    root.after(100, update_video_frame)

def start_video(source):
    global cap, running
    stop_video()
    cap = cv2.VideoCapture(source)
    if not cap.isOpened():
        print("Error: Cannot open video source.")
        return
    running = True
    update_video_frame()

frame_skip_counter = 0
frame_skip_rate = 10

def update_video_frame():
    global cap, running, frame_skip_counter, last_processed_frame
    if not running or cap is None:
        return

    ret, frame = cap.read()
    if not ret:
        print("Error: Failed to read frame")
        stop_video()
        return

    frame = cv2.resize(frame, (640, 360))
    frame_skip_counter += 1
    if frame_skip_counter >= frame_skip_rate:
        frame_skip_counter = 0
        results = model(frame, conf=detection_threshold)
        frame_copy = frame.copy()

        for r in results:
            if r.boxes is not None:
                for box in r.boxes:
                    x1, y1, x2, y2 = map(int, box.xyxy[0])
                    cls_id = int(box.cls[0])
                    conf = float(box.conf[0])
                    label = f"{model.names[cls_id]} ({conf:.2f})"
                    cv2.rectangle(frame_copy, (x1, y1), (x2, y2), (57, 255, 20), 2)
                    cv2.putText(frame_copy, label, (x1, y1 - 10), cv2.FONT_HERSHEY_SIMPLEX, 0.6, (0, 255, 255), 2)

        last_processed_frame = frame_copy
        display_defects(results, img=last_processed_frame)

    show_image(last_processed_frame if last_processed_frame is not None else frame)

    if running:
        root.after(500, update_video_frame)

def stop_video():
    global cap, running
    running = False
    if cap is not None:
        cap.release()
        cap = None
    save_images_to_excel()

def on_close():
    global cap, ser
    if cap is not None:
        cap.release()
    if ser is not None and ser.is_open:
        ser.close()
    root.destroy
    
# -------- GUI Widgets --------

# Detection Threshold Slider
threshold_slider = tk.Scale(control_frame, from_=0.1, to=1.0, resolution=0.05,
                            orient="horizontal", label="Detection Threshold",
                            command=update_threshold)
threshold_slider.set(0.5)
threshold_slider.pack(pady=5, fill='x')

# COM Port Dropdown
tk.Label(control_frame, text="COM Port").pack()
com_port_var = tk.StringVar()
com_ports = list_serial_ports()
if com_ports:
    com_port_var.set(com_ports[0])
com_port_menu = tk.OptionMenu(control_frame, com_port_var, *com_ports)
com_port_menu.pack(fill='x')

# Baud Rate Dropdown
tk.Label(control_frame, text="Baud Rate").pack()
baud_rate_var = tk.StringVar(value="9600")
tk.OptionMenu(control_frame, baud_rate_var, "4800", "9600", "19200", "38400", "57600", "115200").pack(fill='x')

# Serial Buttons
connect_btn = tk.Button(control_frame, text="Connect Serial", command=connect_serial)
connect_btn.pack(pady=5, fill='x')

tk.Button(control_frame, text="Disconnect Serial", command=disconnect_serial).pack(pady=5, fill='x')

# File/Camera Control Buttons
tk.Button(control_frame, text="Open Image", command=open_image).pack(pady=5, fill='x')
tk.Button(control_frame, text="Open Video", command=open_video).pack(pady=5, fill='x')
tk.Button(control_frame, text="Start Camera", command=start_camera).pack(pady=5, fill='x')
tk.Button(control_frame, text="Stop", command=stop_video).pack(pady=5, fill='x')
tk.Button(control_frame, text="Save Images to Excel", command=save_images_to_excel).pack(pady=5, fill='x')

# GPS label
gps_label = tk.Label(control_frame, text="GPS: N/A", fg="black", anchor="w", justify="left", font=("Arial", 10))
gps_label.pack(pady=10, fill='x')

# Final setup and start
root.protocol("WM_DELETE_WINDOW", on_close)
update_gps_label()
root.mainloop()
